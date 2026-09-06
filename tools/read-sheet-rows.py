"""Dump one worksheet as JSON rows, so a Node tool can read a workbook.

The repository already splits this work in two: Python reads .xlsx (openpyxl),
and the tool that writes data/active-state.json is JavaScript, because the keys
it writes come from slot-model.js and reimplementing ownedKey in a second
language is how two spellings of the same card start to drift apart.

    python3 tools/read-sheet-rows.py <workbook.xlsx> [sheet name]

Emits [[cell, ...], ...] including the header row, values only.
"""
import json
import sys

import openpyxl


def main() -> int:
    if len(sys.argv) < 2:
        print(__doc__, file=sys.stderr)
        return 2
    book = openpyxl.load_workbook(sys.argv[1], data_only=True, read_only=True)
    sheet = book[sys.argv[2]] if len(sys.argv) > 2 else book[book.sheetnames[0]]
    rows = [
        [None if cell is None else (cell if isinstance(cell, (int, float)) else str(cell))
         for cell in row]
        for row in sheet.iter_rows(values_only=True)
    ]
    json.dump(rows, sys.stdout)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
